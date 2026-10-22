"""Tạo Báo cáo Kết quả Kinh doanh Q3 2026.xlsx (Lab 1, Lab 3, Lab 7).

Công ty dịch vụ giả định, không nêu tên. Đơn vị: tỷ đồng (dự án: triệu đồng). Số liệu được thiết kế để có các "điểm lệch":
  1) Doanh thu tăng nhưng lợi nhuận giảm vì chi phí thuê thầu phụ vượt kế hoạch
  2) Tỷ lệ nghỉ việc cao hơn ngành, NPS giảm, phụ thuộc vào một khách hàng lớn
Chạy xong in đáp án tham khảo cho Lab 3 Bài tập 2 (phân loại dự án).
"""
import random
from datetime import date, timedelta
from pathlib import Path
from statistics import median
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

OUT = Path(__file__).resolve().parent.parent / "Báo cáo Kết quả Kinh doanh Q3 2026.xlsx"
random.seed(2027)
AS_OF = date(2026, 9, 30)

F = "Arial"
BLUE = Font(name=F, color="0000FF")
BLACK = Font(name=F)
GREEN = Font(name=F, color="008000")
BOLD = Font(name=F, bold=True)
HDR = Font(name=F, bold=True, color="FFFFFF")
HDR_FILL = PatternFill("solid", fgColor="0B3D5C")
SUB_FILL = PatternFill("solid", fgColor="E8EEF3")
TITLE = Font(name=F, bold=True, size=14, color="0B3D5C")
NOTE = Font(name=F, italic=True, size=9, color="595959")
thin = Side(style="thin", color="BFBFBF")
BOX = Border(top=thin, bottom=thin, left=thin, right=thin)
NUM = '#,##0;(#,##0);"-"'
NUM1 = '#,##0.0;(#,##0.0);"-"'
PCT = '0.0%;(0.0%);"-"'

wb = Workbook()


def header(ws, row, cols):
    for i, c in enumerate(cols, 1):
        cell = ws.cell(row=row, column=i, value=c)
        cell.font = HDR
        cell.fill = HDR_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BOX


def widths(ws, ws_widths):
    for i, w in enumerate(ws_widths, 1):
        ws.column_dimensions[get_column_letter(i)].width = w


def title(ws, text, note):
    ws["A1"] = text
    ws["A1"].font = TITLE
    ws["A2"] = note
    ws["A2"].font = NOTE


def add_table(ws, ref, name):
    t = Table(displayName=name, ref=ref)
    t.tableStyleInfo = TableStyleInfo(name="TableStyleLight9", showRowStripes=True)
    ws.add_table(t)


ws_out = wb.active
ws_out.title = "Tổng quan Q3"

# ---------------------------------------------------------------- PL_Budget_Actual
ws = wb.create_sheet("PL_Budget_Actual")
title(ws, "Kết quả hoạt động - Kế hoạch và thực hiện Q3/2026",
      "Đơn vị: tỷ đồng, trừ khi ghi chú khác. Chi phí ghi số dương. Số liệu giả tưởng, chỉ dùng cho mục đích đào tạo.")
header(ws, 4, ["Khoản mục", "Q2/2026 TH", "Q3/2026 KH", "Q3/2026 TH", "TH vs KH", "% vs KH",
               "TH vs Q2", "% vs Q2", "Ghi chú"])
pl = [
    # (key, nhãn, Q2, KH, TH, kiểu, ghi chú)  kiểu: h=tiêu đề nhóm, v=giá trị, s=tổng, s1=tổng 1 số lẻ, p=tỷ lệ (công thức), pv=tỷ lệ (nhập tay)
    (None, "DOANH THU VÀ LỢI NHUẬN", None, None, None, "h", ""),
    ("dt1", "Doanh thu Tư vấn", 120, 125, 121, "v", ""),
    ("dt2", "Doanh thu Triển khai giải pháp", 166, 175, 178, "v", "Ba dự án lớn nghiệm thu sớm trong tháng 9"),
    ("dt3", "Doanh thu Dịch vụ vận hành", 98, 100, 104, "v", ""),
    ("dt4", "Doanh thu Đào tạo doanh nghiệp", 34, 36, 29, "v", "Hai khóa đào tạo lớn bị khách hàng hoãn sang Q4"),
    ("dt", "Tổng doanh thu", "sum:dt1:dt4", None, None, "s", ""),
    ("cp1", "Chi phí nhân sự", 205, 214, 219, "v", "Bao gồm chi phí làm thêm giờ"),
    ("cp2", "Chi phí thuê thầu phụ", 38, 36, 50, "v", "Thuê ngoài để bù nhân sự nghỉ việc ở mảng Triển khai giải pháp"),
    ("cp3", "Chi phí văn phòng và hành chính", 34, 35, 35, "v", ""),
    ("cp4", "Chi phí marketing và bán hàng", 22, 26, 24, "v", ""),
    ("cp5", "Chi phí CNTT và bản quyền phần mềm", 18, 19, 20, "v", ""),
    ("cp", "Tổng chi phí", "sum:cp1:cp5", None, None, "s", ""),
    ("ln", "Lợi nhuận trước thuế", "=dt-cp", None, None, "s", ""),
    ("bien", "Biên lợi nhuận trước thuế", "=ln/dt", None, None, "p", ""),
    (None, "KHÁCH HÀNG", None, None, None, "h", ""),
    ("kh", "Số khách hàng đang phục vụ", 286, 300, 291, "v", ""),
    ("tt", "Tỷ trọng doanh thu từ khách hàng lớn nhất", 0.31, 0.28, 0.36, "pv", "Khách hàng KH-001 (đã ẩn danh)"),
    ("nps", "Điểm hài lòng khách hàng (NPS)", 41, 45, 36, "v", "Khảo sát khách hàng cuối quý"),
    (None, "NHÂN SỰ", None, None, None, "h", ""),
    ("nv", "Số nhân viên cuối kỳ (người)", 1185, 1220, 1162, "v", ""),
    ("nghi", "Số nhân viên nghỉ việc trong quý (người)", 44, 35, 58, "v", "Tập trung ở văn phòng TP. Hồ Chí Minh"),
    ("tlnghi", "Tỷ lệ nghỉ việc năm hóa", "=nghi*4/nv", None, None, "p", "Bằng số nghỉ việc trong quý x 4 chia số nhân viên cuối kỳ"),
    ("ot", "Giờ làm thêm bình quân / người / tháng", 9.5, 8, 14.2, "v1", ""),
    ("dtao", "Giờ đào tạo bình quân / người trong quý", 7.5, 8, 4.8, "v1", ""),
    (None, "CHỈ SỐ HIỆU QUẢ", None, None, None, "h", ""),
    ("dtnv", "Doanh thu bình quân / nhân viên (triệu đồng)", "=dt*1000/nv", None, None, "s1", ""),
    ("cir", "Tỷ lệ chi phí trên doanh thu", "=cp/dt", None, None, "p", ""),
]
ROW = {}
for i, (key, *_rest) in enumerate(pl, start=5):
    if key:
        ROW[key] = i


def expand(expr, col):
    """'sum:dt1:dt4' hoặc '=ln/dt' -> công thức Excel cho cột col."""
    if expr.startswith("sum:"):
        _, a, b = expr.split(":")
        return f"=SUM({col}{ROW[a]}:{col}{ROW[b]})"
    out = expr
    for k in sorted(ROW, key=len, reverse=True):
        out = out.replace(k, f"{col}{ROW[k]}")
    return out


for i, (key, lab, q2, kh, th, kind, note) in enumerate(pl, start=5):
    ws.cell(row=i, column=1, value=lab)
    if kind == "h":
        for col in range(1, 10):
            ws.cell(row=i, column=col).fill = SUB_FILL
        ws.cell(row=i, column=1).font = BOLD
        continue
    if kind in ("s", "s1", "p"):
        q2, kh, th = (expand(q2, c) for c in "BCD")
    is_pct = kind in ("p", "pv")
    fmt = PCT if is_pct else (NUM1 if kind in ("s1", "v1") else NUM)
    for col, v in ((2, q2), (3, kh), (4, th)):
        c = ws.cell(row=i, column=col, value=v)
        c.font = BLUE if kind in ("v", "v1", "pv") else (BOLD if kind == "s" else BLACK)
        c.number_format = fmt
    ws.cell(row=i, column=1).font = BOLD if kind == "s" else BLACK
    if is_pct:  # chênh lệch điểm phần trăm
        ws.cell(row=i, column=5, value=f"=D{i}-C{i}").number_format = PCT
        ws.cell(row=i, column=7, value=f"=D{i}-B{i}").number_format = PCT
        ws.cell(row=i, column=6, value="điểm %")
        ws.cell(row=i, column=8, value="điểm %")
    else:
        ws.cell(row=i, column=5, value=f"=D{i}-C{i}").number_format = fmt
        ws.cell(row=i, column=6, value=f"=IFERROR(D{i}/C{i}-1,0)").number_format = PCT
        ws.cell(row=i, column=7, value=f"=D{i}-B{i}").number_format = fmt
        ws.cell(row=i, column=8, value=f"=IFERROR(D{i}/B{i}-1,0)").number_format = PCT
    for col in range(5, 9):
        ws.cell(row=i, column=col).font = BLACK
    ws.cell(row=i, column=9, value=note).font = NOTE
widths(ws, [44, 13, 13, 13, 12, 10, 12, 10, 58])
ws.freeze_panes = "B5"
PL = "PL_Budget_Actual"

# ---------------------------------------------------------------- Doanh thu tháng
monthly = [  # tháng, DT KH, CP KH, DT TH, CP TH
    (1, 130, 101, 128, 99), (2, 120, 97, 118, 94), (3, 150, 107, 149, 107),
    (4, 138, 106, 136, 104), (5, 141, 107, 139, 105), (6, 146, 109, 143, 108),
    (7, 143, 108, 142, 112), (8, 145, 110, 141, 115), (9, 148, 112, 149, 121),
]
assert sum(m[3] for m in monthly[6:]) == 432 and sum(m[4] for m in monthly[6:]) == 348
assert sum(m[3] for m in monthly[3:6]) == 418 and sum(m[4] for m in monthly[3:6]) == 317
ws = wb.create_sheet("Doanh thu tháng")
title(ws, "Doanh thu và chi phí theo tháng - 9 tháng 2026",
      "Đơn vị: tỷ đồng. KH = kế hoạch, TH = thực hiện.")
header(ws, 4, ["Tháng", "Doanh thu KH", "Chi phí KH", "Lợi nhuận KH", "Doanh thu TH", "Chi phí TH",
               "Lợi nhuận TH", "Chênh lệch lợi nhuận TH - KH", "% doanh thu so với KH"])
for i, (m, dk, ck, dt_, ct) in enumerate(monthly, start=5):
    ws.cell(row=i, column=1, value=f"T{m}/2026").font = BLACK
    for col, v in ((2, dk), (3, ck), (5, dt_), (6, ct)):
        c = ws.cell(row=i, column=col, value=v)
        c.font = BLUE
        c.number_format = NUM
    for col, f, fmt in ((4, f"=B{i}-C{i}", NUM), (7, f"=E{i}-F{i}", NUM),
                        (8, f"=G{i}-D{i}", NUM), (9, f"=IFERROR(E{i}/B{i}-1,0)", PCT)):
        c = ws.cell(row=i, column=col, value=f)
        c.font = BLACK
        c.number_format = fmt
add_table(ws, "A4:I13", "DoanhThuThang")
ws.cell(row=14, column=1, value="Lũy kế 9T").font = BOLD
for col in range(2, 9):
    L = get_column_letter(col)
    c = ws.cell(row=14, column=col, value=f"=SUM({L}5:{L}13)")
    c.font = BOLD
    c.number_format = NUM
c = ws.cell(row=14, column=9, value="=IFERROR(E14/B14-1,0)")
c.font = BOLD
c.number_format = PCT
widths(ws, [12, 13, 12, 13, 13, 12, 13, 18, 14])
ws.freeze_panes = "A5"

# ---------------------------------------------------------------- Danh sách dự án
offices = [
    # (văn phòng, số NV, nghỉ việc Q3, doanh thu Q3, chi phí Q3, NPS, giờ làm thêm, số dự án, hệ số chậm, hệ số vượt chi phí)
    ("Hà Nội", 380, 17, 146, 115, 38, 13.1, 46, 1.0, 1.0),
    ("TP. Hồ Chí Minh", 410, 27, 160, 136, 29, 18.6, 52, 1.5, 1.6),
    ("Đà Nẵng", 130, 4, 46, 35, 47, 9.2, 16, 0.6, 0.6),
    ("Hải Phòng", 82, 3, 28, 22, 40, 10.4, 11, 0.9, 0.8),
    ("Cần Thơ", 70, 3, 22, 18, 42, 11.8, 10, 1.1, 1.2),
    ("Bình Dương", 90, 4, 30, 22, 35, 12.5, 15, 0.8, 0.9),
]
assert sum(o[1] for o in offices) == 1162 and sum(o[2] for o in offices) == 58
assert sum(o[3] for o in offices) == 432 and sum(o[4] for o in offices) == 348
lines = [("Tư vấn", 0.9, 0.8), ("Triển khai giải pháp", 1.3, 1.4), ("Dịch vụ vận hành", 0.7, 0.8), ("Đào tạo doanh nghiệp", 1.0, 0.9)]

projects = []
pm_pool = [f"PM{n:02d}" for n in range(1, 37)]
for name, *_, n_proj, f_late, f_over in offices:
    for _ in range(n_proj):
        line, l_late, l_over = random.choices(lines, weights=[3, 4, 3, 1.2])[0]
        start = AS_OF - timedelta(days=random.randint(40, 330))
        end = AS_OF + timedelta(days=random.randint(20, 260))
        t = (AS_OF - start).days / (end - start).days
        done = max(0.05, min(0.98, t + random.gauss(0.04, 0.08) - random.random() * 0.13 * f_late * l_late))
        cost_ratio = done * (1 + random.gauss(-0.07, 0.08) + random.random() * 0.13 * f_over * l_over)
        budget = random.choice([3, 4, 5, 6, 8, 10, 12, 15, 18, 25, 30, 40, 55, 70]) * 100
        projects.append(dict(
            vp=name, mang=line, pm=random.choice(pm_pool),
            kh="KH-001" if random.random() < 0.18 else f"KH-{random.randint(2, 291):03d}",
            start=start, end=end, budget=budget,
            cost=round(budget * cost_ratio), done=round(done, 2),
            csat=round(min(5, max(2, random.gauss(4.1 - 0.5 * (f_late - 1), 0.45))), 1),
        ))
random.shuffle(projects)
projects.sort(key=lambda p: p["start"])

ws = wb.create_sheet("Danh sách dự án")
title(ws, "Danh sách dự án đang triển khai tại 30/09/2026",
      "Ngân sách và chi phí: triệu đồng. Mã khách hàng và mã quản lý dự án đã ẩn danh. "
      "% thời gian đã qua tính đến 30/09/2026. Số liệu giả tưởng.")
cols = ["Mã dự án", "Văn phòng", "Mảng dịch vụ", "Quản lý dự án", "Khách hàng", "Ngày bắt đầu",
        "Ngày kết thúc KH", "Ngân sách", "Chi phí lũy kế", "% hoàn thành", "% thời gian đã qua", "Điểm hài lòng KH (1-5)"]
header(ws, 4, cols)
ws.row_dimensions[4].height = 32
for i, p in enumerate(projects, start=5):
    p["ma"] = f"DA{i - 4:03d}"
    vals = [p["ma"], p["vp"], p["mang"], p["pm"], p["kh"], p["start"], p["end"], p["budget"], p["cost"], p["done"],
            f"=MIN(1,MAX(0,(DATE(2026,9,30)-F{i})/(G{i}-F{i})))", p["csat"]]
    for j, v in enumerate(vals, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font = BLACK if j <= 5 or j == 11 else BLUE
        if j in (6, 7):
            c.number_format = "dd/mm/yyyy"
        elif j in (8, 9):
            c.number_format = NUM
        elif j in (10, 11):
            c.number_format = "0%"
        elif j == 12:
            c.number_format = "0.0"
last_p = 4 + len(projects)
add_table(ws, f"A4:L{last_p}", "DanhSachDuAn")
widths(ws, [10, 17, 22, 13, 12, 13, 15, 11, 13, 12, 14, 14])
ws.freeze_panes = "B5"
DA = "'Danh sách dự án'"

# ---------------------------------------------------------------- Theo văn phòng
ws = wb.create_sheet("Theo văn phòng")
title(ws, "Hiệu quả theo văn phòng - Q3/2026",
      "Đơn vị: tỷ đồng, trừ khi ghi chú khác. Số dự án lấy công thức từ sheet Danh sách dự án.")
cols = ["Văn phòng", "Số nhân viên", "Nghỉ việc Q3", "Tỷ lệ nghỉ việc năm hóa", "Doanh thu Q3",
        "Doanh thu / nhân viên (triệu đồng)", "Chi phí Q3", "Biên lợi nhuận", "NPS", "Giờ làm thêm / người / tháng",
        "Số dự án đang chạy"]
header(ws, 4, cols)
ws.row_dimensions[4].height = 32
for i, (name, nv, nghi, dt_, cp_, nps, ot, *_r) in enumerate(offices, start=5):
    row = [name, nv, nghi, f"=IFERROR(C{i}*4/B{i},0)", dt_, f"=IFERROR(E{i}*1000/B{i},0)", cp_,
           f"=IFERROR((E{i}-G{i})/E{i},0)", nps, ot, f"=COUNTIF({DA}!$B$5:$B${last_p},$A{i})"]
    fmts = [None, NUM, NUM, PCT, NUM, NUM1, NUM, PCT, NUM, NUM1, NUM]
    for j, (v, fm) in enumerate(zip(row, fmts), 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font = GREEN if j == 11 else (BLACK if isinstance(v, str) or j == 1 else BLUE)
        if fm:
            c.number_format = fm
endo = 4 + len(offices)
add_table(ws, f"A4:K{endo}", "TheoVanPhong")
t = endo + 1
ws.cell(row=t, column=1, value="Toàn công ty").font = BOLD
for col, f, fm in (("B", f"=SUM(B5:B{endo})", NUM), ("C", f"=SUM(C5:C{endo})", NUM),
                   ("D", f"=IFERROR(C{t}*4/B{t},0)", PCT), ("E", f"=SUM(E5:E{endo})", NUM),
                   ("F", f"=IFERROR(E{t}*1000/B{t},0)", NUM1), ("G", f"=SUM(G5:G{endo})", NUM),
                   ("H", f"=IFERROR((E{t}-G{t})/E{t},0)", PCT), ("K", f"=SUM(K5:K{endo})", NUM)):
    c = ws[f"{col}{t}"]
    c.value = f
    c.font = BOLD
    c.number_format = fm
widths(ws, [18, 12, 12, 14, 12, 16, 11, 12, 8, 15, 12])
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Benchmark
ws = wb.create_sheet("Benchmark")
title(ws, "So sánh với nhóm doanh nghiệp dịch vụ cùng quy mô - 9 tháng 2026",
      "Nguồn: tổng hợp nội bộ từ báo cáo công bố và khảo sát ngành (đã ẩn danh). Số liệu giả tưởng, chỉ dùng cho mục đích đào tạo.")
cols = ["Doanh nghiệp", "Doanh thu 9 tháng (tỷ đồng)", "Tăng trưởng doanh thu so với cùng kỳ", "Biên lợi nhuận trước thuế",
        "Tỷ lệ nghỉ việc năm hóa", "NPS", "Giờ đào tạo / người / năm", "Tỷ trọng doanh thu từ khách hàng lớn nhất"]
header(ws, 4, cols)
ws.row_dimensions[4].height = 48
peers = [
    ("Doanh nghiệp cùng ngành A", 1240, 0.11, 0.23, 0.12, 44, 36, 0.18),
    ("Doanh nghiệp cùng ngành B", 990, 0.09, 0.21, 0.135, 40, 30, 0.22),
    ("Doanh nghiệp cùng ngành C", 1580, 0.14, 0.26, 0.15, 48, 42, 0.15),
    ("Doanh nghiệp cùng ngành D", 735, 0.06, 0.18, 0.16, 35, 24, 0.27),
    ("Doanh nghiệp cùng ngành E", 1110, 0.10, 0.22, 0.13, 46, 32, 0.20),
]
FMT_B = [None, NUM, PCT, PCT, PCT, NUM, NUM, PCT]
for i, row in enumerate(peers, start=5):
    for j, v in enumerate(row, 1):
        c = ws.cell(row=i, column=j, value=v)
        c.font = BLUE if j >= 2 else BLACK
        if FMT_B[j - 1]:
            c.number_format = FMT_B[j - 1]
orow = 5 + len(peers)
ours = ["Công ty của chúng ta", "='Doanh thu tháng'!E14", 0.08, f"={PL}!D{ROW['bien']}", f"={PL}!D{ROW['tlnghi']}",
        f"={PL}!D{ROW['nps']}", f"={PL}!D{ROW['dtao']}*4", f"={PL}!D{ROW['tt']}"]
for j, v in enumerate(ours, 1):
    c = ws.cell(row=orow, column=j, value=v)
    c.font = Font(name=F, bold=True, color="008000") if isinstance(v, str) and j > 1 else (
        Font(name=F, bold=True, color="0000FF") if j > 1 else BOLD)
    c.fill = PatternFill("solid", fgColor="FFF2CC")
    if FMT_B[j - 1]:
        c.number_format = FMT_B[j - 1]
add_table(ws, f"A4:H{orow}", "Benchmark")
med = orow + 1
ws.cell(row=med, column=1, value="Trung vị nhóm so sánh (A-E)").font = BOLD
for col in range(2, 9):
    L = get_column_letter(col)
    c = ws.cell(row=med, column=col, value=f"=MEDIAN({L}5:{L}{orow - 1})")
    c.font = BOLD
    c.number_format = FMT_B[col - 1]
ws.cell(row=med + 1, column=1, value="Chênh lệch so với trung vị").font = BOLD
for col in range(3, 9):
    L = get_column_letter(col)
    c = ws.cell(row=med + 1, column=col, value=f"={L}{orow}-{L}{med}")
    c.number_format = FMT_B[col - 1]
ws.cell(row=med + 3, column=1,
        value="Ghi chú: giờ đào tạo của Công ty năm hóa từ số liệu Q3; tăng trưởng doanh thu so với cùng kỳ (8%) do Phòng Tài chính cung cấp.").font = NOTE
widths(ws, [30, 15, 16, 14, 14, 8, 14, 18])
ws.freeze_panes = "B5"

# ---------------------------------------------------------------- Tổng quan Q3
ws = ws_out
title(ws, "Tổng quan kết quả Q3/2026",
      "Đơn vị: tỷ đồng, trừ khi ghi chú khác. Tất cả số liệu liên kết từ sheet PL_Budget_Actual.")
header(ws, 4, ["Chỉ tiêu", "Q2/2026 TH", "Q3/2026 KH", "Q3/2026 TH", "% hoàn thành KH", "Tăng trưởng so với Q2"])
kpis = [  # (nhãn, key, kiểu)
    ("Tổng doanh thu", "dt", NUM), ("Tổng chi phí", "cp", NUM), ("Lợi nhuận trước thuế", "ln", NUM),
    ("Biên lợi nhuận trước thuế", "bien", PCT), ("Số khách hàng đang phục vụ", "kh", NUM),
    ("Điểm hài lòng khách hàng (NPS)", "nps", NUM), ("Số nhân viên cuối kỳ", "nv", NUM),
    ("Tỷ lệ nghỉ việc năm hóa", "tlnghi", PCT),
]
for i, (lab, key, fmt) in enumerate(kpis, start=5):
    ws.cell(row=i, column=1, value=lab).font = BLACK
    for col, src in ((2, "B"), (3, "C"), (4, "D")):
        c = ws.cell(row=i, column=col, value=f"={PL}!{src}{ROW[key]}")
        c.font = GREEN
        c.number_format = fmt
    if fmt == PCT:
        ws.cell(row=i, column=5, value="n/a").alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=6, value=f"=D{i}-B{i}").number_format = '+0.0%;-0.0%;"-"'
    elif key == "nps":
        ws.cell(row=i, column=5, value="n/a").alignment = Alignment(horizontal="right")
        ws.cell(row=i, column=6, value=f"=D{i}-B{i}").number_format = '+0;-0;"-"'
    else:
        ws.cell(row=i, column=5, value=f"=IFERROR(D{i}/C{i},0)").number_format = PCT
        ws.cell(row=i, column=6, value=f"=IFERROR(D{i}/B{i}-1,0)").number_format = PCT
ws.cell(row=14, column=1, value="Ghi chú: với chỉ tiêu tỷ lệ và NPS, cột tăng trưởng thể hiện chênh lệch điểm.").font = NOTE
widths(ws, [34, 13, 13, 13, 16, 20])

for s in wb.worksheets:
    for row in s.iter_rows():
        for c in row:
            if c.font and c.font.name != F:
                c.font = Font(name=F, bold=c.font.bold, italic=c.font.italic, color=c.font.color, size=c.font.size)
wb.calculation.fullCalcOnLoad = True
wb.save(OUT)
print("saved", OUT, f"({len(projects)} dự án)")

# ---------------------------------------------------------------- Đáp án tham khảo Lab 3 BT2
counts = {o[0]: {"Đúng hướng": 0, "Cần chú ý": 0, "Rủi ro": 0} for o in offices}
for p in projects:
    t = min(1, max(0, (AS_OF - p["start"]).days / (p["end"] - p["start"]).days))
    late = p["done"] < t - 0.10
    over = p["cost"] / p["budget"] > p["done"]
    status = "Rủi ro" if late and over else ("Đúng hướng" if not late and not over else "Cần chú ý")
    counts[p["vp"]][status] += 1
tot = {k: sum(c[k] for c in counts.values()) for k in ("Đúng hướng", "Cần chú ý", "Rủi ro")}
print("Lab 3 BT2 - toàn công ty:", tot)
for vp, c in sorted(counts.items(), key=lambda kv: -kv[1]["Rủi ro"] / sum(kv[1].values())):
    n = sum(c.values())
    print(f"  {vp:16s} {c}  tỷ lệ Rủi ro {c['Rủi ro'] / n:.0%}")
print("KH-001:", sum(p["kh"] == "KH-001" for p in projects), "dự án")
print("Trung vị nghỉ việc ngành:", median(p[4] for p in peers))
