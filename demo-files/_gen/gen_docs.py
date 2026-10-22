"""Tạo các file Word demo cho workshop (bản general, công ty giả định không nêu tên)."""
import random
from pathlib import Path
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

from docx_common import new_doc, title, para, bullets, table, footer_note, set_font

OUT = Path(__file__).resolve().parent.parent
FICTION = "Tài liệu nháp (draft) dùng cho mục đích đào tạo. Tên nhân vật, khách hàng, đối tác và số liệu là giả định."


# ============================================================ Lab 2: Dự thảo hợp đồng dịch vụ
def contract():
    doc = new_doc()
    footer_note(doc, FICTION)
    title(doc, "HỢP ĐỒNG CUNG CẤP DỊCH VỤ TỔ CHỨC SỰ KIỆN",
          "Số: 1026/2026/HĐDV-LK  |  Dự thảo v0.2 ngày 12/10/2026  |  Lưu hành nội bộ")
    para(doc, "Hợp đồng này (\"Hợp đồng\") được lập giữa:")
    bullets(doc, [
        ("Bên A - Bên sử dụng dịch vụ (\"Công ty\"): ", "trụ sở tại TP. Hồ Chí Minh; "
         "đại diện bởi Bà Lê Thu Hà, Giám đốc Hành chính - Nhân sự."),
        ("Bên B - Công ty TNHH Sự kiện Lam Kiều (\"Lam Kiều\"): ", "trụ sở tại TP. Thủ Đức, TP. Hồ Chí Minh; "
         "đại diện bởi Ông Phan Đức Minh, Giám đốc."),
    ])

    doc.add_heading("Điều 1. Định nghĩa", 1)
    bullets(doc, [
        ("1.1. Sự kiện: ", "gồm Sự kiện 1 - Tiệc tất niên 2026 cho khoảng 650 nhân viên, tổ chức ngày 16/01/2027 tại TP. Hồ Chí Minh; "
         "và Sự kiện 2 - Hội nghị Khách hàng 2027 cho khoảng 250 khách mời, tổ chức ngày 12/03/2027 tại Hà Nội."),
        ("1.2. Kịch bản: ", "tài liệu mô tả chương trình, thời lượng, nhân sự, âm thanh ánh sáng, thực đơn và trang trí của từng Sự kiện, "
         "được hai bên ký xác nhận trước ngày tổ chức ít nhất 20 ngày."),
        ("1.3. Nghiệm thu: ", "việc Bên A xác nhận bằng văn bản Bên B đã hoàn thành Sự kiện theo Kịch bản."),
        ("1.4. Chi phí phát sinh: ", "chi phí ngoài báo giá đã duyệt, phát sinh do thay đổi số lượng khách, thực đơn hoặc hạng mục."),
    ], style="Normal")

    doc.add_heading("Điều 2. Phạm vi dịch vụ", 1)
    bullets(doc, [
        "2.1. Tư vấn ý tưởng, xây dựng Kịch bản và thiết kế nhận diện cho từng Sự kiện.",
        "2.2. Thuê địa điểm, âm thanh, ánh sáng, màn hình LED, MC, ca sĩ và tiết mục biểu diễn.",
        "2.3. Tiệc buffet (Sự kiện 1) và tiệc trà, ăn trưa (Sự kiện 2) theo thực đơn được Bên A duyệt.",
        "2.4. Gửi thư mời điện tử, quản lý đăng ký và check-in khách mời; tổ chức bốc thăm trúng thưởng tại Sự kiện 1.",
        "2.5. Chụp ảnh, quay phim và bàn giao tư liệu trong 07 ngày sau mỗi Sự kiện.",
    ], style="Normal")

    doc.add_heading("Điều 3. Giá trị hợp đồng", 1)
    table(doc, ["Hạng mục", "Giá trị (chưa VAT)"], [
        ["Sự kiện 1 - Tiệc tất niên 2026 (650 khách)", "1.480.000.000 đồng"],
        ["Sự kiện 2 - Hội nghị Khách hàng 2027 (250 khách)", "870.000.000 đồng"],
        ["Tổng giá trị hợp đồng", "2.350.000.000 đồng"],
    ], [10, 6])
    para(doc, "Giá trị trên đã bao gồm toàn bộ chi phí nhân sự, vận chuyển, thiết bị và thuế, phí liên quan (trừ thuế GTGT).")

    doc.add_heading("Điều 4. Quyền và nghĩa vụ của Bên B", 1)
    bullets(doc, [
        "4.1. Thực hiện Sự kiện đúng Kịch bản, đúng thời gian, bảo đảm an toàn cho người tham dự.",
        "4.2. Bố trí 01 quản lý dự án làm đầu mối duy nhất, có mặt tại các buổi họp chuẩn bị theo yêu cầu của Bên A.",
        "4.3. Mua bảo hiểm cho nhân sự và thiết bị của Bên B trong thời gian thực hiện Sự kiện.",
        "4.4. Chịu trách nhiệm với nhà cung cấp phụ do Bên B thuê.",
    ], style="Normal")

    doc.add_heading("Điều 5. Quyền và nghĩa vụ của Bên A", 1)
    bullets(doc, [
        "5.1. Cung cấp thông tin khách mời, nhận diện thương hiệu và duyệt Kịch bản đúng hạn.",
        "5.2. Cử 01 đầu mối phối hợp và phê duyệt thay đổi trong quá trình chuẩn bị.",
        "5.3. Thanh toán theo Điều 8 sau khi nhận đủ hồ sơ thanh toán hợp lệ.",
    ], style="Normal")

    doc.add_heading("Điều 6. Bảo mật thông tin", 1)
    para(doc, "6.1. Bên B bảo mật mọi thông tin nhận được từ Bên A, chỉ sử dụng cho mục đích thực hiện Hợp đồng và không chuyển cho bên thứ ba.")
    para(doc, "6.2. Việc xử lý dữ liệu cá nhân của nhân viên và khách mời của Bên A tuân thủ pháp luật về bảo vệ dữ liệu cá nhân.")

    doc.add_heading("Điều 7. Nghiệm thu", 1)
    para(doc, "Trong 03 ngày làm việc sau mỗi Sự kiện, hai bên ký biên bản nghiệm thu. Bên A có quyền ghi nhận các hạng mục không đạt "
         "để làm căn cứ giảm trừ thanh toán.")

    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(14)
    r = p.add_run("GHI CHÚ CỦA NGƯỜI SOẠN THẢO - Nội dung đã thống nhất tại buổi làm việc ngày 08/10/2026, cần đưa vào Điều 8 và Điều 9:")
    r.bold = True
    r.font.color.rgb = RGBColor(0xC0, 0x50, 0x00)
    bullets(doc, [
        "Thanh toán: tạm ứng 50% giá trị hợp đồng (1.175.000.000 đồng) trong 03 ngày làm việc sau khi ký; phần còn lại thanh toán trong 05 ngày làm việc sau mỗi Sự kiện.",
        "Chi phí phát sinh: Bên B được chủ động quyết định chi phí phát sinh đến 10% giá trị mỗi Sự kiện, Bên A thanh toán theo thực tế.",
        "Hủy, đổi lịch: Bên A hủy trước ngày tổ chức từ 30 ngày trở lên chịu 50% giá trị Sự kiện; hủy trong vòng 15 ngày chịu 100%.",
        "Phạt vi phạm: nếu Bên B chậm tiến độ hoặc thực hiện sai Kịch bản, Bên B chịu phạt 5% giá trị Sự kiện đó. Mức bồi thường thiệt hại tối đa chưa thống nhất.",
        "Dữ liệu khách mời: Bên A gửi danh sách 650 nhân viên gồm họ tên, số điện thoại, ngày sinh và email cá nhân để Bên B gửi thư mời và tổ chức bốc thăm. "
        "Bên B được sử dụng hình ảnh, video Sự kiện cho mục đích giới thiệu năng lực của Bên B.",
        "Báo giá: đã nhận 02 báo giá (Lam Kiều 2,35 tỷ đồng; một đơn vị khác 2,60 tỷ đồng). Đơn vị thứ ba chưa gửi báo giá.",
        "Chưa chốt: bảo hiểm trách nhiệm công cộng cho Sự kiện 1 (tổ chức ngoài trời) và phương án dự phòng khi thời tiết xấu.",
    ], style="Normal")
    doc.save(OUT / "Hợp đồng Dịch vụ Tổ chức Sự kiện (dự thảo).docx")


# ============================================================ Lab 2 mở rộng: Chính sách bảo mật (bản lỗi định dạng)
def broken_policy():
    random.seed(11)
    doc = new_doc()
    footer_note(doc, FICTION)
    fonts = ["Times New Roman", "Calibri", "Courier New", "Cambria", "VNI-Times", ".VnTime", "Arial Narrow"]
    tone_map = str.maketrans("ảãạáàẩẫậấầẳẵặắằẻẽẹéèểễệếềỉĩịíìỏõọóòổỗộốồởỡợớờủũụúùửữựứừỷỹỵýỳ",
                             "".join(ch * 5 for ch in "aâăeêioôơuưy"))

    def corrupt(text):
        out = []
        for w in text.split(" "):
            x = random.random()
            if x < 0.10:
                w = w.translate(tone_map)  # mất dấu thanh
            elif x < 0.13 and len(w) > 4:
                k = random.randint(1, len(w) - 2)
                w = w[:k] + " " + w[k:]  # tách chữ do chuyển đổi PDF
            out.append(w)
        return " ".join(out)

    def line(text, heading=False):
        # mô phỏng PDF -> Word: ngắt dòng giữa câu, font và cỡ chữ lộn xộn
        chunks = [text]
        if not heading and len(text) > 110 and random.random() < 0.6:
            cut = text.rfind(" ", 0, random.randint(60, 100))
            chunks = [text[:cut], text[cut + 1:]]
        for ch in chunks:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(random.choice([0, 0, 2, 9, 14]))
            p.alignment = random.choice([WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.JUSTIFY,
                                         WD_ALIGN_PARAGRAPH.CENTER if heading else WD_ALIGN_PARAGRAPH.LEFT])
            p.paragraph_format.left_indent = Pt(random.choice([0, 0, 6, 18, 30]))
            r = p.add_run(corrupt(ch))
            set_font(r, random.choice(fonts))
            r.font.size = Pt(random.choice([12, 13, 14, 16]) if heading else random.choice([9, 10, 10.5, 11, 12, 13]))
            r.bold = heading or random.random() < 0.05
            r.italic = random.random() < 0.06

    line("TÀI LIỆU NỘI BỘ", True)
    line("CHÍNH SÁCH BẢO MẬT THÔNG TIN VÀ SỬ DỤNG TRÍ TUỆ NHÂN TẠO", True)
    line("(Ban hành kèm Quyết định số 0815/2026/QĐ-TGĐ ngày 15/08/2026)  Trang 1/4")
    sections = [
        ("I. MỤC ĐÍCH", [
            "Chính sách này quy định nguyên tắc, trách nhiệm và biện pháp bảo vệ thông tin của Công ty, khách hàng, đối tác và nhân viên, "
            "bảo đảm tuân thủ pháp luật về bảo vệ dữ liệu cá nhân và các cam kết bảo mật trong hợp đồng với khách hàng.",
        ]),
        ("2 . Phạm vi áp dụng", [
            "Áp dụng cho toàn bộ nhân viên, cộng tác viên, thực tập sinh và nhà cung cấp có quyền truy cập hệ thống hoặc tài liệu của Công ty, "
            "ở mọi phòng ban, văn phòng và khi làm việc từ xa.",
        ]),
        ("III- Định nghĩa", [
            "Thông tin nội bộ: mọi thông tin phát sinh trong quá trình hoạt động của Công ty chưa được công bố chính thức.",
            "Dữ liệu cá nhân: họ tên, ngày sinh, số CCCD hoặc hộ chiếu, số điện thoại, email cá nhân, địa chỉ, số tài khoản ngân hàng, hình ảnh.",
            "Dữ liệu cá nhân nhạy cảm: thông tin lương, thưởng, sức khỏe, đánh giá hiệu quả công việc và kỷ luật của nhân viên.",
        ]),
        ("4) Phân loại thông tin", [
            "Cấp 1 - Công khai: thông tin đã công bố trên website, tài liệu giới thiệu dịch vụ.",
            "Cấp 2 - Nội bộ: quy trình, hướng dẫn, biên bản họp, báo cáo tổng hợp không chứa dữ liệu cá nhân.",
            "Cấp 3 - Mật: hợp đồng, báo giá, báo cáo tài chính chi tiết, dữ liệu khách hàng, dữ liệu cá nhân của nhân viên.",
            "Cấp 4 - Tuyệt mật: kế hoạch sáp nhập, thông tin lương thưởng cá nhân, mật khẩu, khóa mã hóa và thông tin xác thực hệ thống.",
        ]),
        ("V. NGUYÊN TẮC BẢO MẬT", [
            "Chỉ thu thập dữ liệu cá nhân cần thiết cho mục đích đã thông báo và được chủ thể dữ liệu đồng ý.",
            "Phân quyền truy cập theo nguyên tắc cần biết (need-to-know); thư mục dùng chung phải có người chịu trách nhiệm phân quyền.",
            "Không lưu thông tin Cấp 3, Cấp 4 trên thiết bị cá nhân, email cá nhân hay ứng dụng nhắn tin không được phê duyệt.",
            "Mọi tài liệu Cấp 3, Cấp 4 phải được gắn nhãn phân loại (sensitivity label) trước khi chia sẻ.",
        ]),
        ("6. Cung cấp thông tin cho bên thứ ba", [
            "Chỉ cung cấp thông tin tối thiểu cần thiết theo hợp đồng đã ký. Nhà cung cấp xử lý dữ liệu cá nhân của nhân viên hoặc khách hàng "
            "phải ký điều khoản xử lý dữ liệu cá nhân; việc dùng hình ảnh nhân viên cho mục đích quảng bá phải được nhân viên đồng ý.",
        ]),
        ("VII. SỬ DỤNG CÔNG CỤ TRÍ TUỆ NHÂN TẠO", [
            "7.1. Công cụ được phép: chỉ sử dụng Microsoft 365 Copilot, Copilot Chat và các agent chạy trong môi trường Microsoft 365 của Công ty. "
            "Không đưa dữ liệu nội bộ vào công cụ AI công cộng hoặc tài khoản AI cá nhân.",
            "7.2. Dữ liệu được phép: thông tin Cấp 1, Cấp 2 được sử dụng bình thường. Thông tin Cấp 3 chỉ được xử lý bằng Copilot khi tài liệu "
            "đã gắn nhãn phân loại và người dùng có quyền truy cập hợp lệ. Thông tin Cấp 4 không được đưa vào bất kỳ công cụ AI nào.",
            "7.3. Không nhập số CCCD, hộ chiếu, số tài khoản ngân hàng, mật khẩu, thông tin lương cá nhân vào câu lệnh (prompt); "
            "khi cần nói về một người hay một khách hàng cụ thể, dùng mã đã ẩn danh.",
            "7.4. Con người chịu trách nhiệm cuối cùng: AI không được tự ra quyết định tuyển dụng, đánh giá, kỷ luật nhân viên hay phê duyệt chi tiêu. "
            "Nội dung do AI tạo phải được người dùng kiểm tra số liệu và câu chữ trước khi gửi đi.",
            "7.5. Ghi âm, ghi lời thoại (transcript) cuộc họp có khách hàng hoặc đối tác bên ngoài chỉ được bật sau khi thông báo và được "
            "các bên đồng ý ngay đầu cuộc họp. Biên bản do AI tạo phải được người chủ trì rà soát trước khi gửi.",
            "7.6. Agent tự tạo chỉ được chia sẻ trong nội bộ Công ty; knowledge source phải là thư mục SharePoint có phân quyền. "
            "Agent phục vụ khách hàng bên ngoài phải được Phòng Công nghệ Thông tin và Phòng Pháp chế phê duyệt.",
            "7.7. Kết quả do AI tạo được lưu trên OneDrive hoặc SharePoint của Công ty và gắn nhãn phân loại như tài liệu gốc.",
            "7.8. Báo cáo cho Bộ phận An ninh thông tin trong vòng 24 giờ khi phát hiện AI trả về thông tin sai lệch nghiêm trọng "
            "hoặc hiển thị dữ liệu mà người dùng không có quyền xem.",
        ]),
        ("8 - Xử lý sự cố lộ lọt thông tin", [
            "Người phát hiện sự cố báo cáo ngay cho Bộ phận An ninh thông tin (security@company.example) và quản lý trực tiếp trong vòng 02 giờ.",
            "Phòng Pháp chế đánh giá mức độ ảnh hưởng, thông báo cho các bên liên quan và cơ quan quản lý theo thời hạn luật định.",
        ]),
        ("IX. TRÁCH NHIỆM THI HÀNH", [
            "Trưởng các phòng ban chịu trách nhiệm phổ biến chính sách. Vi phạm bị xử lý kỷ luật theo Nội quy lao động và "
            "có thể bị truy cứu trách nhiệm theo quy định pháp luật.",
        ]),
    ]
    for h, paras in sections:
        line(h, True)
        for t in paras:
            line(("- " if random.random() < 0.5 else "• ") + t if len(paras) > 1 else t)
        if random.random() < 0.4:
            line(f"Trang {random.randint(1, 4)}/4")
    doc.save(OUT / "SAMPLE Chính sách Bảo mật Thông tin.docx")


# ============================================================ Lab 1 + Lab 4: Kế hoạch hoạt động 2027
def plan():
    doc = new_doc()
    footer_note(doc, FICTION)
    title(doc, "KẾ HOẠCH HOẠT ĐỘNG NĂM 2027",
          "Dự thảo trình Ban Giám đốc  |  Đơn vị soạn thảo: Phòng Kế hoạch và Chiến lược  |  Phiên bản 0.4 ngày 05/10/2026")

    doc.add_heading("1. Bối cảnh", 1)
    para(doc, "Sau 9 tháng 2026, doanh thu đạt 1.245 tỷ đồng, bằng 98,7% kế hoạch 9 tháng; dự kiến cả năm đạt khoảng 1.690 tỷ đồng. "
         "Tỷ lệ nghỉ việc năm hóa tăng lên 20%, cao hơn đáng kể so với nhóm doanh nghiệp cùng ngành (khoảng 13,5%). "
         "Điểm hài lòng khách hàng (NPS) giảm từ 41 xuống 36. Khảo sát nhân viên quý 3 cho thấy ba lý do chính khiến nhân viên muốn nghỉ: "
         "không thấy lộ trình phát triển, khối lượng làm thêm giờ lớn và thiếu thời gian đào tạo.")

    doc.add_heading("2. Mục tiêu đến 31/12/2027", 1)
    table(doc, ["Chỉ tiêu", "Ước 2026", "Mục tiêu 2027"], [
        ["Doanh thu (tỷ đồng)", "1.690", "1.900"],
        ["Tỷ lệ nghỉ việc năm hóa", "20%", "Dưới 13%"],
        ["Điểm hài lòng khách hàng (NPS)", "+36", "+45"],
        ["Giờ đào tạo bình quân / người / năm", "19", "32"],
        ["Giờ làm thêm bình quân / người / tháng", "14", "Dưới 8"],
        ["Tỷ lệ nhân viên dùng Microsoft 365 Copilot hằng tuần", "Chưa triển khai", "80%"],
    ], [8.5, 3.5, 3.5])

    doc.add_heading("3. Bốn sáng kiến trọng tâm", 1)
    bullets(doc, [
        ("Sáng kiến 1 - Con người: ", "lộ trình nghề nghiệp rõ ràng cho từng vị trí, chương trình mentor, tăng ngân sách đào tạo, "
         "đánh giá 1:1 hằng tháng giữa quản lý và nhân viên."),
        ("Sáng kiến 2 - Khách hàng: ", "phản hồi mọi yêu cầu của khách hàng trong 24 giờ làm việc, khảo sát sau mỗi giai đoạn dự án, "
         "họp đánh giá chất lượng dịch vụ hằng quý với 30 khách hàng lớn nhất."),
        ("Sáng kiến 3 - Làm việc thông minh với AI: ", "triển khai Microsoft 365 Copilot cho toàn bộ nhân viên, đào tạo theo vai trò, "
         "xây dựng agent nội bộ tra cứu chính sách và quy trình."),
        ("Sáng kiến 4 - Tăng trưởng: ", "mở văn phòng Huế trong quý 3/2027, tăng doanh thu mảng Đào tạo doanh nghiệp thêm 30%."),
    ], style="Normal")

    doc.add_heading("4. Chương trình Phát triển và Phúc lợi 2027", 1)
    para(doc, "Áp dụng cho nhân viên chính thức từ ngày 01/01/2027, thay thế các mức tương ứng trong Sổ tay Chính sách Nội bộ hiện hành.")
    table(doc, ["Quyền lợi", "Nhân viên", "Chuyên viên, Trưởng nhóm", "Quản lý"], [
        ["Ngân sách đào tạo / năm", "8 triệu đồng", "12 triệu đồng", "25 triệu đồng"],
        ["Làm việc từ xa", "Tối đa 2 ngày/tuần", "Tối đa 2 ngày/tuần", "Linh hoạt theo lịch công tác"],
        ["Ngày nghỉ thêm", "Nghỉ sinh nhật 1 ngày", "Nghỉ sinh nhật 1 ngày", "Nghỉ sinh nhật 1 ngày"],
        ["Khám sức khỏe hằng năm", "Gói cơ bản", "Gói nâng cao", "Gói nâng cao, thêm 1 người thân"],
        ["Mentor", "Được ghép mentor trong 6 tháng đầu", "Làm mentor, phụ cấp 1 triệu đồng/tháng", "Làm mentor cho trưởng nhóm"],
        ["Thưởng giới thiệu ứng viên", "8 triệu đồng", "8 triệu đồng", "15 triệu đồng"],
    ], [4.5, 3.8, 4, 4])

    doc.add_heading("5. Vai trò của quản lý và từng nhân viên", 1)
    bullets(doc, [
        ("Tháng 1/2027: ", "mỗi quản lý trao đổi 1:1 với từng thành viên để thống nhất lộ trình nghề nghiệp và kế hoạch đào tạo cá nhân."),
        ("Hằng tháng: ", "quản lý rà soát giờ làm thêm của đội, điều chỉnh phân công khi một người vượt 12 giờ/tháng."),
        ("Hằng quý: ", "mỗi nhân viên hoàn thành tối thiểu 8 giờ đào tạo, trong đó có 2 giờ về sử dụng AI an toàn và hiệu quả."),
        ("Tuân thủ: ", "chỉ sử dụng công cụ AI được Công ty phê duyệt; không đưa dữ liệu cá nhân hay dữ liệu khách hàng vào công cụ AI công cộng."),
    ], style="Normal")

    doc.add_heading("6. Lộ trình triển khai", 1)
    table(doc, ["Thời gian", "Hoạt động", "Đơn vị chủ trì"], [
        ["11/2026", "Phê duyệt kế hoạch, công bố khung lộ trình nghề nghiệp", "Phòng Nhân sự"],
        ["12/2026", "Cấp licence Microsoft 365 Copilot, đào tạo nhóm Copilot Champion", "Phòng Công nghệ Thông tin"],
        ["01/2027", "Áp dụng Chương trình Phát triển và Phúc lợi 2027, trao đổi 1:1", "Phòng Nhân sự, các phòng ban"],
        ["Quý 1/2027", "Đào tạo Copilot theo vai trò cho toàn bộ nhân viên", "Phòng Công nghệ Thông tin"],
        ["Quý 3/2027", "Mở văn phòng Huế", "Ban Giám đốc"],
        ["Hằng quý", "Đánh giá tiến độ, điều chỉnh kế hoạch", "Phòng Kế hoạch và Chiến lược"],
    ], [3, 8, 5])

    doc.add_heading("7. Chỉ tiêu theo văn phòng đến 31/12/2027", 1)
    table(doc, ["Văn phòng", "Nhân viên hiện tại", "Mục tiêu doanh thu 2027 (tỷ đồng)", "Mục tiêu tỷ lệ nghỉ việc"], [
        ["Hà Nội", "380", "640", "Dưới 13%"],
        ["TP. Hồ Chí Minh", "410", "700", "Dưới 15%"],
        ["Đà Nẵng", "130", "200", "Dưới 12%"],
        ["Hải Phòng, Cần Thơ, Bình Dương", "242", "310", "Dưới 13%"],
        ["Văn phòng Huế (mở mới)", "-", "50", "-"],
        ["Toàn công ty", "1.162", "1.900", "Dưới 13%"],
    ], [5, 3.5, 4.5, 3.5])

    doc.add_heading("8. Ngân sách dự kiến", 1)
    para(doc, "Tổng chi phí thực hiện kế hoạch năm 2027 dự kiến 64 tỷ đồng, gồm đào tạo và mentor (18 tỷ), phúc lợi tăng thêm (26 tỷ), "
         "licence Copilot và hạ tầng CNTT (12 tỷ), truyền thông nội bộ và tuyển dụng (8 tỷ). Lợi ích ước tính 150 tỷ đồng nhờ giảm chi phí "
         "tuyển dụng thay thế, tăng năng suất và giữ chân khách hàng.")
    doc.save(OUT / "Kế hoạch Hoạt động 2027 (dự thảo).docx")


# ============================================================ Lab 2 + Lab 8: Sổ tay chính sách (knowledge source)
def handbook():
    doc = new_doc()
    footer_note(doc, FICTION + " Phiên bản 2.1 - hiệu lực từ 01/07/2026.")
    title(doc, "SỔ TAY CHÍNH SÁCH NỘI BỘ",
          "Dành cho toàn bộ nhân viên  |  Phiên bản 2.1, hiệu lực từ 01/07/2026  |  Đơn vị quản lý: Phòng Hành chính - Nhân sự")

    doc.add_heading("Mục 1. Hướng dẫn sử dụng Sổ tay", 1)
    para(doc, "Sổ tay tổng hợp các chính sách nội bộ thường dùng: thời gian làm việc, nghỉ phép, làm việc từ xa, công tác phí, mua sắm, "
         "thẩm quyền phê duyệt, đào tạo, bảo mật và phúc lợi. Trường hợp có khác biệt giữa Sổ tay và văn bản gốc, áp dụng văn bản gốc. "
         "Câu hỏi ngoài phạm vi Sổ tay, nhân viên liên hệ Phòng Hành chính - Nhân sự qua hộp thư hr.policy@company.example.")

    doc.add_heading("Mục 2. Thời gian làm việc và nghỉ phép", 1)
    doc.add_heading("2.1. Thời gian làm việc", 2)
    bullets(doc, [
        "Từ thứ Hai đến thứ Sáu, 8 giờ 30 đến 17 giờ 30, nghỉ trưa 1 giờ. Nhân viên được linh hoạt giờ bắt đầu từ 7 giờ 30 đến 9 giờ 30, bảo đảm đủ 8 giờ/ngày.",
        "Làm thêm giờ phải được quản lý trực tiếp phê duyệt trước trên hệ thống; tối đa 40 giờ/tháng theo quy định pháp luật.",
    ])
    doc.add_heading("2.2. Nghỉ phép năm", 2)
    table(doc, ["Thâm niên làm việc tại Công ty", "Số ngày phép năm"], [
        ["Dưới 5 năm", "14 ngày (12 ngày theo luật + 2 ngày Công ty tặng thêm)"],
        ["Từ 5 năm đến dưới 10 năm", "15 ngày"],
        ["Từ 10 năm trở lên", "16 ngày"],
    ], [8, 8])
    bullets(doc, [
        "Đăng ký nghỉ phép trên hệ thống trước ít nhất 03 ngày làm việc; nghỉ liên tục từ 05 ngày trở lên phải đăng ký trước 02 tuần.",
        "Ngày phép chưa dùng hết được chuyển sang năm sau tối đa 05 ngày và phải sử dụng trước ngày 31/03; quá hạn không được thanh toán.",
        "Nghỉ ốm: nộp giấy xác nhận của cơ sở y tế trong 03 ngày làm việc sau khi đi làm lại.",
    ])

    doc.add_heading("Mục 3. Làm việc từ xa", 1)
    bullets(doc, [
        "Nhân viên đã qua thời gian thử việc được làm việc từ xa tối đa 02 ngày/tuần, do quản lý trực tiếp phê duyệt.",
        "Khi làm việc từ xa phải dùng laptop do Công ty cấp, kết nối VPN; không xử lý tài liệu Mật ở nơi công cộng.",
        "Nhân viên phải có mặt trực tuyến trong khung giờ 9 giờ 30 đến 16 giờ và tham gia đầy đủ các cuộc họp đã lên lịch.",
    ])

    doc.add_heading("Mục 4. Công tác phí trong nước", 1)
    table(doc, ["Hạng mục", "Nhân viên, Chuyên viên", "Trưởng, Phó phòng", "Ban Giám đốc"], [
        ["Khách sạn, tối đa mỗi đêm", "1,2 triệu đồng (Hà Nội, TP. HCM: 1,5 triệu)", "1,8 triệu đồng", "3 triệu đồng"],
        ["Phụ cấp ăn, mỗi ngày công tác", "250.000 đồng", "350.000 đồng", "500.000 đồng"],
        ["Vé máy bay", "Hạng phổ thông", "Hạng phổ thông", "Hạng thương gia nếu chuyến bay trên 3 giờ"],
        ["Di chuyển nội thành", "Taxi, xe công nghệ theo thực tế", "Taxi, xe công nghệ theo thực tế", "Xe Công ty"],
    ], [4, 4.5, 3.5, 4])
    bullets(doc, [
        "Chuyến công tác phải được quản lý trực tiếp phê duyệt trên hệ thống trước khi đặt vé.",
        "Quyết toán trong 10 ngày làm việc sau khi kết thúc chuyến đi, kèm hóa đơn hợp lệ ghi tên Công ty.",
        "Chi phí vượt định mức do nhân viên tự chi trả, trừ khi được Giám đốc khối phê duyệt trước bằng văn bản.",
    ])

    doc.add_heading("Mục 5. Mua sắm, hợp đồng và thẩm quyền phê duyệt", 1)
    doc.add_heading("5.1. Thẩm quyền phê duyệt theo giá trị", 2)
    table(doc, ["Giá trị mua sắm, hợp đồng (chưa VAT)", "Cấp phê duyệt", "Số báo giá tối thiểu"], [
        ["Đến 50 triệu đồng", "Trưởng phòng", "01"],
        ["Trên 50 triệu đến 500 triệu đồng", "Giám đốc khối", "02"],
        ["Trên 500 triệu đến 2 tỷ đồng", "Giám đốc Tài chính", "03"],
        ["Trên 2 tỷ đồng", "Tổng Giám đốc, sau khi Phòng Pháp chế rà soát hợp đồng", "03"],
    ], [6, 7, 3])
    doc.add_heading("5.2. Tạm ứng và thanh toán cho nhà cung cấp", 2)
    bullets(doc, [
        "Tạm ứng tối đa 30% giá trị hợp đồng. Tạm ứng trên 30% là ngoại lệ, cần Giám đốc Tài chính phê duyệt và nhà cung cấp có bảo lãnh tạm ứng của ngân hàng.",
        "Thanh toán trong 30 ngày kể từ ngày nghiệm thu và nhận đủ hóa đơn hợp lệ. Thanh toán sớm hơn cần Kế toán trưởng đồng ý.",
        "Không tách nhỏ một nhu cầu mua sắm thành nhiều hợp đồng để tránh thẩm quyền phê duyệt.",
    ])
    doc.add_heading("5.3. Điều khoản bắt buộc trong hợp đồng với nhà cung cấp", 2)
    bullets(doc, [
        "Mọi chi phí phát sinh ngoài báo giá phải được Công ty chấp thuận bằng văn bản trước khi thực hiện.",
        "Phí hủy hoặc đổi lịch trên 30% giá trị hợp đồng phải được Phòng Pháp chế rà soát trước khi ký.",
        "Có điều khoản phạt vi phạm khi nhà cung cấp chậm tiến độ hoặc không đạt chất lượng; mức phạt đề xuất 8% giá trị phần nghĩa vụ bị vi phạm "
        "(mức tối đa theo Luật Thương mại).",
        "Nhà cung cấp xử lý dữ liệu cá nhân của nhân viên hoặc khách hàng phải ký điều khoản xử lý dữ liệu cá nhân theo Mục 7.3.",
    ])

    doc.add_heading("Mục 6. Đào tạo và phát triển", 1)
    table(doc, ["Cấp bậc", "Ngân sách đào tạo / người / năm"], [
        ["Nhân viên, Chuyên viên", "6 triệu đồng"],
        ["Trưởng nhóm, Trưởng phòng", "10 triệu đồng"],
        ["Ban Giám đốc", "20 triệu đồng"],
    ], [8, 8])
    bullets(doc, [
        "Đăng ký khóa học bên ngoài trên hệ thống, được quản lý trực tiếp và Phòng Nhân sự phê duyệt.",
        "Khóa học có chi phí trên 30 triệu đồng: nhân viên cam kết làm việc tối thiểu 12 tháng sau khi hoàn thành.",
        "Chứng chỉ chuyên môn liên quan đến công việc được hoàn 100% lệ phí thi nếu đạt.",
    ])

    doc.add_heading("Mục 7. Bảo mật thông tin và sử dụng AI", 1)
    doc.add_heading("7.1. Phân loại thông tin", 2)
    bullets(doc, [
        "Công khai - Nội bộ - Mật - Tuyệt mật. Tài liệu Mật và Tuyệt mật phải gắn nhãn phân loại trước khi chia sẻ.",
        "Hợp đồng, báo giá, dữ liệu khách hàng và dữ liệu cá nhân của nhân viên thuộc loại Mật.",
    ])
    doc.add_heading("7.2. Sử dụng AI", 2)
    bullets(doc, [
        "Chỉ dùng Microsoft 365 Copilot, Copilot Chat và agent trong môi trường Microsoft 365 của Công ty.",
        "Không nhập số CCCD, hộ chiếu, số tài khoản ngân hàng, mật khẩu, thông tin lương cá nhân vào prompt.",
        "Nội dung do AI tạo phải được người dùng kiểm tra trước khi gửi; AI không thay người có thẩm quyền ra quyết định.",
    ])
    doc.add_heading("7.3. Chia sẻ dữ liệu cá nhân cho nhà cung cấp", 2)
    bullets(doc, [
        "Chỉ chia sẻ dữ liệu tối thiểu cần thiết cho mục đích của hợp đồng (ví dụ: họ tên và email công việc để gửi thư mời).",
        "Không chia sẻ ngày sinh, số điện thoại cá nhân, email cá nhân khi không thật sự cần thiết; nếu cần phải có sự đồng ý của nhân viên.",
        "Nhà cung cấp chỉ được dùng hình ảnh của nhân viên cho mục đích quảng bá khi có sự đồng ý bằng văn bản của Công ty và nhân viên.",
    ])

    doc.add_heading("Mục 8. Thiết bị và hỗ trợ CNTT", 1)
    bullets(doc, [
        "Laptop được cấp mới sau 04 năm sử dụng hoặc khi hỏng không sửa được.",
        "Mất thiết bị hoặc nghi ngờ lộ mật khẩu: báo Phòng Công nghệ Thông tin qua hotline nội bộ 1900 trong vòng 02 giờ.",
        "Phần mềm ngoài danh mục chuẩn phải được Phòng Công nghệ Thông tin phê duyệt trước khi cài đặt.",
    ])

    doc.add_heading("Mục 9. Phúc lợi", 1)
    bullets(doc, [
        "Khám sức khỏe định kỳ hằng năm và bảo hiểm sức khỏe cho nhân viên chính thức.",
        "Thưởng giới thiệu ứng viên: 5 triệu đồng cho vị trí nhân viên, 10 triệu đồng cho vị trí quản lý, chi trả khi ứng viên qua thử việc.",
        "Quà sinh nhật, hỗ trợ hiếu hỉ và hoạt động gắn kết theo quy chế của Công đoàn.",
    ])

    doc.add_heading("Mục 10. Đầu mối liên hệ", 1)
    table(doc, ["Nội dung", "Đơn vị"], [
        ["Nghỉ phép, làm việc từ xa, phúc lợi, đào tạo", "Phòng Hành chính - Nhân sự"],
        ["Công tác phí, tạm ứng, thanh toán", "Phòng Tài chính - Kế toán"],
        ["Hợp đồng, mua sắm", "Phòng Pháp chế, Bộ phận Mua sắm"],
        ["Thiết bị, tài khoản, bảo mật", "Phòng Công nghệ Thông tin"],
    ], [8, 8])
    doc.save(OUT / "Sổ tay Chính sách Nội bộ.docx")


if __name__ == "__main__":
    contract()
    broken_policy()
    plan()
    handbook()
    print("done")
